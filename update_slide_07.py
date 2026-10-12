import os
import subprocess
from pptx import Presentation
from pptx.util import Pt

pptx_path = r"C:\Users\sharv\barcode-test\Provenance_Client_Presentation.pptx"

# Ensure PowerPoint is closed
subprocess.run(["powershell", "-Command", "Stop-Process -Name POWERPNT -Force -ErrorAction SilentlyContinue"], capture_output=True)

prs = Presentation(pptx_path)
slide = prs.slides[6] # Slide 7 (0-indexed: 6)

def update_run_text(shape, paragraph_idx, text, font_size_pt=None):
    p = shape.text_frame.paragraphs[paragraph_idx]
    if p.runs:
        p.runs[0].text = text
        if font_size_pt:
            p.runs[0].font.size = Pt(font_size_pt)
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.text = text
        if font_size_pt:
            p.font.size = Pt(font_size_pt)

def set_paragraphs(shape, lines, font_size_pt=11.5):
    tf = shape.text_frame
    tf.word_wrap = True
    # Clean up excess paragraphs if any
    while len(tf.paragraphs) > len(lines):
        p_elem = tf.paragraphs[-1]._p
        p_elem.getparent().remove(p_elem)
        
    for i, line in enumerate(lines):
        if i < len(tf.paragraphs):
            p = tf.paragraphs[i]
            p.space_after = Pt(3)
            p.space_before = Pt(0)
            if p.runs:
                p.runs[0].text = line
                p.runs[0].font.size = Pt(font_size_pt)
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.text = line
                p.font.size = Pt(font_size_pt)
        else:
            p = tf.add_paragraph()
            p.text = line
            p.font.size = Pt(font_size_pt)
            p.space_after = Pt(3)
            p.space_before = Pt(0)

# Card 1: Camera
update_run_text(slide.shapes[12], 0, "01 • BENCH CAMERA")
update_run_text(slide.shapes[13], 0, "1080p", font_size_pt=46.0)
card1_bullets = [
    "1080p Full HD industrial sensor",
    "Adaptive lighting tuning",
    "Anti-glare tape optimization",
    "Wide bench angle coverage",
    "Rigid industrial mount arm"
]
set_paragraphs(slide.shapes[14], card1_bullets, font_size_pt=11.5)

# Card 2: Barcode Scanner
update_run_text(slide.shapes[18], 0, "02 • BARCODE SCANNER")
update_run_text(slide.shapes[19], 0, "IP65", font_size_pt=46.0)
card2_bullets = [
    "IP65 industrial casing",
    "USB and wireless support",
    "1D and 2D barcode ready",
    "Fast reads on torn labels",
    "Zero keyboard typing"
]
set_paragraphs(slide.shapes[20], card2_bullets, font_size_pt=11.5)

# Card 3: Station PC
update_run_text(slide.shapes[24], 0, "03 • STATION PC")
update_run_text(slide.shapes[25], 0, "16GB", font_size_pt=46.0)
card3_bullets = [
    "Dust-sealed fanless chassis",
    "Quad-core edge processor",
    "16GB RAM visual buffer",
    "Fast local solid state drive",
    "Auto 90-day clip storage"
]
set_paragraphs(slide.shapes[26], card3_bullets, font_size_pt=11.5)

# Card 4: Bench Feedback
update_run_text(slide.shapes[30], 0, "04 • BENCH FEEDBACK")
update_run_text(slide.shapes[31], 0, "85dB", font_size_pt=46.0)
card4_bullets = [
    "Live on-screen indicator",
    "Audible start/stop tones",
    "Heard over warehouse noise",
    "Desk-mounted finish card",
    "Zero worker ambiguity"
]
set_paragraphs(slide.shapes[32], card4_bullets, font_size_pt=11.5)

prs.save(pptx_path)
print("Slide 7 updated with tight fitting bullets")
