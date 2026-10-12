import os
import subprocess
from pptx import Presentation
from pptx.util import Pt

pptx_path = r"C:\Users\sharv\barcode-test\Provenance_Client_Presentation.pptx"

# Ensure PowerPoint is closed
subprocess.run(["powershell", "-Command", "Stop-Process -Name POWERPNT -Force -ErrorAction SilentlyContinue"], capture_output=True)

prs = Presentation(pptx_path)
slide = prs.slides[8] # Slide 9 (0-indexed: 8)

def update_run_text(shape, paragraph_idx, text, font_size_pt=None):
    p = shape.text_frame.paragraphs[paragraph_idx]
    if p.runs:
        p.runs[0].text = text
        if font_size_pt:
            p.runs[0].font.size = Pt(font_size_pt)
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.text = text
        if font_size_pt:
            p.font.size = Pt(font_size_pt)

def set_card_bullets(shape, bullets, font_size_pt=12.5):
    tf = shape.text_frame
    tf.word_wrap = True
    while len(tf.paragraphs) > len(bullets):
        p_elem = tf.paragraphs[-1]._p
        p_elem.getparent().remove(p_elem)
        
    for i, bullet in enumerate(bullets):
        if i < len(tf.paragraphs):
            p = tf.paragraphs[i]
            p.space_after = Pt(4)
            p.space_before = Pt(0)
            if p.runs:
                p.runs[0].text = bullet
                p.runs[0].font.size = Pt(font_size_pt)
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.text = bullet
                p.font.size = Pt(font_size_pt)
        else:
            p = tf.add_paragraph()
            p.text = bullet
            p.font.size = Pt(font_size_pt)
            p.space_after = Pt(4)
            p.space_before = Pt(0)

# Card 1 (Week 01)
update_run_text(slide.shapes[14], 0, "Bench Audit & Setup")
card1_bullets = [
    "Industrial camera bench mount",
    "Handheld barcode integration",
    "Field-of-view & lighting tuning",
    "Zero-clutter cabling & power setup"
]
set_card_bullets(slide.shapes[15], card1_bullets, font_size_pt=12.5)

# Card 2 (Week 02)
update_run_text(slide.shapes[21], 0, "Station Pilot Launch")
card2_bullets = [
    "Single packing bench activation",
    "5s visual buffer validation",
    "Worker zero-friction observation",
    "90-day local clip storage check"
]
set_card_bullets(slide.shapes[22], card2_bullets, font_size_pt=12.5)

# Card 3 (Week 03)
update_run_text(slide.shapes[28], 0, "Full Facility Expansion")
card3_bullets = [
    "Multi-station facility rollout",
    "Central network storage backup",
    "Claims team portal onboarding",
    "Dispute defense workflow dry-run"
]
set_card_bullets(slide.shapes[29], card3_bullets, font_size_pt=12.5)

# Card 4 (Week 04)
update_run_text(slide.shapes[34], 0, "WMS & Courier Integration")
card4_bullets = [
    "WMS barcode metadata linking",
    "One-click courier video export",
    "Executive ROI dashboard live",
    "Full production sign-off"
]
set_card_bullets(slide.shapes[35], card4_bullets, font_size_pt=12.5)

prs.save(pptx_path)
print("Slide 9 updated successfully in Provenance_Client_Presentation.pptx")
