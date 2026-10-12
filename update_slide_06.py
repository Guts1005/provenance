import os
import subprocess
from pptx import Presentation

pptx_path = r"C:\Users\sharv\barcode-test\Provenance_Client_Presentation.pptx"

# Ensure PowerPoint is closed
subprocess.run(["powershell", "-Command", "Stop-Process -Name POWERPNT -Force -ErrorAction SilentlyContinue"], capture_output=True)

prs = Presentation(pptx_path)
slide = prs.slides[5] # Slide 6 (0-indexed: 5)

def update_run_text(shape, paragraph_idx, text):
    p = shape.text_frame.paragraphs[paragraph_idx]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.text = text

# Header & Subtitle
update_run_text(slide.shapes[7], 0, "Technical Architecture")
update_run_text(slide.shapes[7], 1, "Reliable Local-First Architecture")
update_run_text(slide.shapes[8], 0, "Engineered to operate uninterrupted on local hardware, immune to warehouse Wi-Fi drops.")

# Tier 03 (Claims Portal)
update_run_text(slide.shapes[15], 0, "Dispute & Claims Portal")
update_run_text(slide.shapes[16], 0, "Instant barcode search by AWB or order ID. Generates a compact MP4 video proof clip ready for one-click submission to courier and marketplace dispute portals.")
update_run_text(slide.shapes[17], 0, "Sub-2s")
update_run_text(slide.shapes[18], 0, "BARCODE SEARCH")

# Tier 02 (Edge Engine)
update_run_text(slide.shapes[23], 0, "EDGE ENGINE")
update_run_text(slide.shapes[25], 0, "Station Edge Engine")
update_run_text(slide.shapes[26], 0, "Local-first background service that records, buffers, and indexes video locally on the bench PC. Runs 24/7 with zero dependency on warehouse cloud or network uptime.")
update_run_text(slide.shapes[27], 0, "100%")
update_run_text(slide.shapes[28], 0, "OFFLINE RESILIENT")

# Tier 01 (Bench Rig)
update_run_text(slide.shapes[34], 0, "Hardware Bench Rig")
update_run_text(slide.shapes[35], 0, "Industrial 1080p camera with dynamic lighting tuning, existing handheld barcode scanner, and clear audio-visual packing confirmation.")
update_run_text(slide.shapes[36], 0, "1080p")
update_run_text(slide.shapes[37], 0, "BENCH OPTIMIZED")

prs.save(pptx_path)
print("Slide 6 updated successfully in Provenance_Client_Presentation.pptx")
