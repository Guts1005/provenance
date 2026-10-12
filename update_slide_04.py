import os
import subprocess
from pptx import Presentation
from pptx.util import Pt

pptx_path = r"C:\Users\sharv\barcode-test\Provenance_Client_Presentation.pptx"

# Ensure PowerPoint is closed
subprocess.run(["powershell", "-Command", "Stop-Process -Name POWERPNT -Force -ErrorAction SilentlyContinue"], capture_output=True)

prs = Presentation(pptx_path)
slide = prs.slides[3] # Slide 4 (0-indexed: 3)

def update_run_text(shape, paragraph_idx, text, font_size_pt=None):
    p = shape.text_frame.paragraphs[paragraph_idx]
    if p.runs:
        p.runs[0].text = text
        if font_size_pt:
            p.runs[0].font.size = Pt(font_size_pt)
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.text = text
        if font_size_pt:
            p.font.size = Pt(font_size_pt)

# Title & Subtitle
update_run_text(slide.shapes[7], 0, "The Solution")
update_run_text(slide.shapes[7], 1, "Automated Packing Bench Verification")
update_run_text(slide.shapes[8], 0, "Zero-friction station software and high-definition video capture designed specifically for fast-paced warehouse benches.")

# Subsystem 01
update_run_text(slide.shapes[12], 0, "0 Clicks", font_size_pt=34.0)
update_run_text(slide.shapes[13], 0, "ZERO WORKER FRICTION")
update_run_text(slide.shapes[14], 0, "SUBSYSTEM 01")
update_run_text(slide.shapes[15], 0, "Instant 0-Click Barcode Trigger")
update_run_text(slide.shapes[16], 0, "Seamlessly intercepts barcode scans from existing USB or wireless handheld scanners. Triggers recording with zero mouse clicks and zero operator slowdown.")

# Subsystem 02
update_run_text(slide.shapes[20], 0, "5 Sec", font_size_pt=34.0)
update_run_text(slide.shapes[21], 0, "CONTINUOUS BUFFER")
update_run_text(slide.shapes[22], 0, "SUBSYSTEM 02")
update_run_text(slide.shapes[23], 0, "5-Second Continuous Visual Buffer")
update_run_text(slide.shapes[24], 0, "Maintains a rolling visual buffer so the clip captures the moments before the scan, preserving complete footage of item placement and box sealing.")

# Subsystem 03
update_run_text(slide.shapes[28], 0, "1080p", font_size_pt=34.0)
update_run_text(slide.shapes[29], 0, "DYNAMIC ADAPTIVE VISION")
update_run_text(slide.shapes[30], 0, "SUBSYSTEM 03")
update_run_text(slide.shapes[31], 0, "1080p Industrial Camera with Lighting Tuning")
update_run_text(slide.shapes[32], 0, "High-definition bench camera automatically adjusts exposure and contrast to eliminate tape glare and maintain sharp package clarity in any warehouse lighting.")

prs.save(pptx_path)
print("Slide 4 updated successfully in Provenance_Client_Presentation.pptx")
