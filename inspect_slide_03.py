from pptx import Presentation

prs = Presentation('C:/Users/sharv/barcode-test/Provenance_Client_Presentation.pptx')
slide = prs.slides[2]
indices = [7, 8, 11, 12, 15, 16, 17, 18, 21, 22, 23, 24, 27, 28, 29, 30, 33, 34, 35, 36, 39, 40, 41, 42, 43, 44, 45]
for idx in indices:
    s = slide.shapes[idx]
    tf = s.text_frame
    parts = []
    for p in tf.paragraphs:
        p_txt = []
        for r in p.runs:
            sz = r.font.size.pt if r.font.size else 'def'
            col = str(r.font.color.rgb) if (r.font.color and r.font.color.type == 1) else 'col'
            p_txt.append(f'[{r.text} | sz={sz} | b={r.font.bold} | {col}]')
        parts.append(' '.join(p_txt))
    print(f"Shape [{idx}] {s.name}: {' || '.join(parts)}")
