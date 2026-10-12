import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

PPTX_PATH = r"C:\Users\sharv\barcode-test\Provenance_Client_Presentation.pptx"

def update_slide_2():
    prs = pptx.Presentation(PPTX_PATH)
    
    # Clean up layout footers globally across the presentation
    for l in prs.slide_layouts:
        for s in l.shapes:
            if s.has_text_frame:
                for p in s.text_frame.paragraphs:
                    for r in p.runs:
                        if "V1.0.4 // PRODUCTION READY" in r.text:
                            r.text = "PROVENANCE STATION SUITE"
                        elif "/ 10" in r.text:
                            r.text = "of 10"

    # Slide 2 (index 1)
    s2 = prs.slides[1]

    # Headline
    s2.shapes[7].text_frame.paragraphs[1].runs[0].text = "Unverifiable Packing Disputes"

    # Subtitle
    s2.shapes[8].text_frame.paragraphs[0].runs[0].text = "When customer shortage or damage claims arise, generic ceiling cameras offer zero proof, forcing merchants to absorb 100% of the loss."

    # Top Metric 1: Rs 3,500 to Rs 15,000
    s2.shapes[9].text_frame.paragraphs[0].runs[0].text = "Rs 3,500 to Rs 15,000"
    s2.shapes[10].text_frame.paragraphs[0].runs[0].text = "AVERAGE LOSS PER CONTESTED ORDER"

    # Top Metric 2: 82% (kept)

    # Timeline Node 1: Stage 01
    s2.shapes[16].text_frame.paragraphs[0].runs[0].text = "STAGE 01 • DAY 0"
    s2.shapes[17].text_frame.paragraphs[0].runs[0].text = "Blind Dispatch"
    s2.shapes[18].text_frame.paragraphs[0].runs[0].text = "Operator scans shipping label and packs order. 20-foot ceiling camera is blocked by worker shoulders."
    s2.shapes[19].text_frame.paragraphs[0].runs[0].text = "Zero visual record of package contents at seal time."

    # Timeline Node 2: Stage 02
    s2.shapes[21].text_frame.paragraphs[0].runs[0].text = "STAGE 02 • DAY 5 TO 10"
    s2.shapes[22].text_frame.paragraphs[0].runs[0].text = "Customer Dispute"
    s2.shapes[23].text_frame.paragraphs[0].runs[0].text = "Buyer claims an empty box or missing item. Marketplace or payment gateway freezes funds automatically."
    s2.shapes[24].text_frame.paragraphs[0].runs[0].text = "Merchant carries 100% burden of proof."

    # Timeline Node 3: Stage 03
    s2.shapes[26].text_frame.paragraphs[0].runs[0].text = "STAGE 03 • DAY 12 TO 18"
    s2.shapes[27].text_frame.paragraphs[0].runs[0].text = "The CCTV Scrub"
    s2.shapes[28].text_frame.paragraphs[0].runs[0].text = "Staff spends 28 minutes searching generic ceiling footage. Tape glare and blur make contents unidentifiable."
    s2.shapes[29].text_frame.paragraphs[0].runs[0].text = "Courier rejects blurry wide-angle clip."

    # Timeline Node 4: Stage 04
    s2.shapes[31].text_frame.paragraphs[0].runs[0].text = "STAGE 04 • DAY 21+"
    s2.shapes[32].text_frame.paragraphs[0].runs[0].text = "Claim Forfeited"
    s2.shapes[33].text_frame.paragraphs[0].runs[0].text = "Courier dispute window closes. Merchant absorbs replacement merchandise, reverse logistics, and penalty charges."
    s2.shapes[34].text_frame.paragraphs[0].runs[0].text = "100% of order value permanently written off."

    prs.save(PPTX_PATH)
    print("Slide 2 updated successfully in:", PPTX_PATH)

    # Export preview
    try:
        import win32com.client
        ppt = win32com.client.Dispatch("PowerPoint.Application")
        pres = ppt.Presentations.Open(PPTX_PATH)
        out_img = r"C:\Users\sharv\.gemini\antigravity\brain\f5a0e924-79bb-495f-9aab-a17922b8c28c\uploaded_preview\slide_02_updated.png"
        pres.Slides(2).Export(out_img, "PNG", 1920, 1080)
        pres.Close()
        ppt.Quit()
        print("Slide 2 preview exported to:", out_img)
    except Exception as e:
        print("Export note:", e)

if __name__ == "__main__":
    update_slide_2()
