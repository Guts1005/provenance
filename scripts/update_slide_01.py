import os
import shutil
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

SRC_PPTX = r"C:\Users\sharv\.gemini\antigravity\brain\f5a0e924-79bb-495f-9aab-a17922b8c28c\.user_uploaded\media_1791732084174_c4031288.pptx"
DEST_PPTX = r"C:\Users\sharv\barcode-test\Provenance_Client_Presentation.pptx"
LOGO_PATH = r"C:\Users\sharv\barcode-test\ui\provenance_logo.png"

def update_slide_1():
    shutil.copyfile(SRC_PPTX, DEST_PPTX)
    prs = pptx.Presentation(DEST_PPTX)
    s1 = prs.slides[0]

    for idx, s in enumerate(s1.shapes):
        if not s.has_text_frame:
            continue
        text = s.text.strip()
        
        # 1. Top Kicker
        if "INDUSTRIAL CCTV VERIFICATION SUITE" in text:
            tf = s.text_frame
            p = tf.paragraphs[0]
            p.text = "PROVENANCE   •   AUTOMATED PACKING SURVEILLANCE"
            p.font.name = "Consolas"
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = RGBColor(16, 185, 129)
            
            s.left = 594360 + Inches(0.44)
            
            if os.path.exists(LOGO_PATH):
                s1.shapes.add_picture(LOGO_PATH, 594360, 830000, width=Inches(0.32), height=Inches(0.32))

        # 2. Main Headline (Clean 2-line title at Pt(36))
        elif "Industrial CCTV" in text and "Packaging" in text:
            s.width = Inches(7.0)
            tf = s.text_frame
            tf.clear()
            tf.word_wrap = True
            
            p1 = tf.paragraphs[0]
            p1.text = "Automated CCTV Evidence &"
            p1.font.name = "Segoe UI"
            p1.font.size = Pt(36)
            p1.font.bold = True
            p1.font.color.rgb = RGBColor(255, 255, 255)
            
            p2 = tf.add_paragraph()
            p2.text = "Packing Station Verification"
            p2.font.name = "Segoe UI"
            p2.font.size = Pt(36)
            p2.font.bold = True
            p2.font.color.rgb = RGBColor(255, 255, 255)

        # 3. Subtitle
        elif "Automated sub-50ms barcode triggered" in text:
            tf = s.text_frame
            p = tf.paragraphs[0]
            p.text = "Instant camera proof for every parcel leaving your facility. Records automatically from your existing barcode scan with zero worker delay."
            p.font.name = "Segoe UI"
            p.font.size = Pt(13)
            p.font.color.rgb = RGBColor(156, 163, 175)

        # 4. Stat Card 1 (Single line metric at Pt(26))
        elif text == "< 50ms":
            tf = s.text_frame
            tf.word_wrap = False
            p = tf.paragraphs[0]
            p.text = "1080p FHD"
            p.font.name = "Segoe UI"
            p.font.size = Pt(26)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)
            
        elif text == "TRIGGER" and s.top > 5500000:
            p = s.text_frame.paragraphs[0]
            p.text = "ADAPTIVE VISION"
            p.font.name = "Consolas"
            p.font.size = Pt(9.5)
            p.font.bold = True
            p.font.color.rgb = RGBColor(16, 185, 129)

        # 5. Stat Card 2
        elif text == "100%":
            tf = s.text_frame
            p = tf.paragraphs[0]
            p.text = "0 Clicks"
            p.font.name = "Segoe UI"
            p.font.size = Pt(26)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)
            
        elif text == "ADMISSIBILITY":
            p = s.text_frame.paragraphs[0]
            p.text = "ZERO WORKER FRICTION"
            p.font.name = "Consolas"
            p.font.size = Pt(9.5)
            p.font.bold = True
            p.font.color.rgb = RGBColor(59, 130, 246)

        # 6. Stat Card 3
        elif text == "82%" and s.left > 5000000:
            tf = s.text_frame
            p = tf.paragraphs[0]
            p.text = "< 2s"
            p.font.name = "Segoe UI"
            p.font.size = Pt(26)
            p.font.bold = True
            p.font.color.rgb = RGBColor(16, 185, 129)
            
        elif text == "DISPUTE WIN RATE" and s.left > 5000000:
            p = s.text_frame.paragraphs[0]
            p.text = "INSTANT EVIDENCE"
            p.font.name = "Consolas"
            p.font.size = Pt(9.5)
            p.font.bold = True
            p.font.color.rgb = RGBColor(16, 185, 129)

        # 7. Camera HUD Header
        elif text == "4K 60FPS":
            p = s.text_frame.paragraphs[0]
            p.text = "1080p FHD"
            p.font.name = "Consolas"
            p.font.size = Pt(10)
            p.font.color.rgb = RGBColor(156, 163, 175)

        # 8. HUD Key-Value Specs
        elif "3840x2160" in text:
            p = s.text_frame.paragraphs[0]
            p.text = "1080p Full HD Industrial Sensor"
            p.font.name = "Consolas"
            p.font.size = Pt(9.5)
            p.font.color.rgb = RGBColor(255, 255, 255)

        elif text == "PRE-ROLL" and s.left > 7000000:
            p = s.text_frame.paragraphs[0]
            p.text = "LIGHTING"
            p.font.name = "Consolas"
            p.font.size = Pt(9.5)
            p.font.color.rgb = RGBColor(107, 114, 128)

        elif "5 S UNCOMPRESSED RAM RING" in text:
            p = s.text_frame.paragraphs[0]
            p.text = "Dynamic anti-glare optimization"
            p.font.name = "Consolas"
            p.font.size = Pt(9.5)
            p.font.color.rgb = RGBColor(156, 163, 175)

        elif text == "TRIGGER" and s.left > 7000000 and s.top > 4000000:
            p = s.text_frame.paragraphs[0]
            p.text = "PRE-ROLL"
            p.font.name = "Consolas"
            p.font.size = Pt(9.5)
            p.font.color.rgb = RGBColor(107, 114, 128)

        elif "< 50 MS KEYSTROKE BURST" in text:
            p = s.text_frame.paragraphs[0]
            p.text = "5-second continuous buffer"
            p.font.name = "Consolas"
            p.font.size = Pt(9.5)
            p.font.color.rgb = RGBColor(156, 163, 175)

        elif text == "STAMP":
            p = s.text_frame.paragraphs[0]
            p.text = "TRIGGER"
            p.font.name = "Consolas"
            p.font.size = Pt(9.5)
            p.font.color.rgb = RGBColor(107, 114, 128)

        elif "SHA-256 + NTP HARDWARE LOCK" in text:
            p = s.text_frame.paragraphs[0]
            p.text = "Standard barcode scan"
            p.font.name = "Consolas"
            p.font.size = Pt(9.5)
            p.font.color.rgb = RGBColor(156, 163, 175)

        elif text == "STATUS":
            p = s.text_frame.paragraphs[0]
            p.text = "STORAGE"
            p.font.name = "Consolas"
            p.font.size = Pt(9.5)
            p.font.color.rgb = RGBColor(107, 114, 128)

        elif "READY // LOCAL-FIRST DEPLOYMENT" in text or "READY" in text:
            p = s.text_frame.paragraphs[0]
            p.text = "Secure local station database"
            p.font.name = "Consolas"
            p.font.size = Pt(9.5)
            p.font.color.rgb = RGBColor(16, 185, 129)

    prs.save(DEST_PPTX)
    print("Slide 1 updated successfully with perfect typography.")

    try:
        import win32com.client
        ppt = win32com.client.Dispatch("PowerPoint.Application")
        pres = ppt.Presentations.Open(DEST_PPTX)
        out_img = r"C:\Users\sharv\.gemini\antigravity\brain\f5a0e924-79bb-495f-9aab-a17922b8c28c\uploaded_preview\slide_01_updated.png"
        pres.Slides(1).Export(out_img, "PNG", 1920, 1080)
        pres.Close()
        ppt.Quit()
        print("Slide 1 preview exported to:", out_img)
    except Exception as e:
        print("Export note:", e)

if __name__ == "__main__":
    update_slide_1()
