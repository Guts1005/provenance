import os
import pptx

SRC_PPTX = r"C:\Users\sharv\.gemini\antigravity\brain\f5a0e924-79bb-495f-9aab-a17922b8c28c\.user_uploaded\media_1791732084174_c4031288.pptx"

def inspect():
    prs = pptx.Presentation(SRC_PPTX)
    print(f"Presentation has {len(prs.slides)} slides.")
    for idx, slide in enumerate(prs.slides):
        print(f"\n--- SLIDE {idx+1} ({len(slide.shapes)} shapes) ---")
        for s_idx, shape in enumerate(slide.shapes):
            if shape.has_text_frame:
                txt = " | ".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
                if txt:
                    print(f"  Shape {s_idx:02d} [{shape.name}]: {txt[:100]}")

if __name__ == "__main__":
    inspect()
