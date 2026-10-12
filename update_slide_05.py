import os
import subprocess
from pptx import Presentation

pptx_path = r"C:\Users\sharv\barcode-test\Provenance_Client_Presentation.pptx"

# Ensure PowerPoint is closed
subprocess.run(["powershell", "-Command", "Stop-Process -Name POWERPNT -Force -ErrorAction SilentlyContinue"], capture_output=True)

prs = Presentation(pptx_path)
slide = prs.slides[4] # Slide 5 (0-indexed: 4)

def update_run_text(shape, paragraph_idx, text):
    p = shape.text_frame.paragraphs[paragraph_idx]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.text = text

# Step 01 description (Shape 14)
update_run_text(slide.shapes[14], 0, "Operator scans shipping label or product barcode with existing handheld scanner. System starts recording automatically with pre-roll footage.")

# Step 02 description (Shape 23)
update_run_text(slide.shapes[23], 0, "Bench camera records continuously in 1080p Full HD while the operator packs. Real-time visual indicator confirms active recording.")

# Step 03 description (Shape 32)
update_run_text(slide.shapes[32], 0, "Operator scans physical desk finish placard. Recording finalizes, saves, and indexes instantly by tracking barcode.")

# Bottom Callout (Shapes 44 & 45)
update_run_text(slide.shapes[44], 0, "BENCH FINISH PLACARD • DURABLE DESK MOUNT")
update_run_text(slide.shapes[45], 0, "A durable barcode placard sits on the packing table within easy reach. One quick scan finalizes the recording and readies the station for the next order.")

prs.save(pptx_path)
print("Slide 5 updated successfully in Provenance_Client_Presentation.pptx")
