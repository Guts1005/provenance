import os
import subprocess
from pptx import Presentation

pptx_path = r"C:\Users\sharv\barcode-test\Provenance_Client_Presentation.pptx"

# Ensure PowerPoint is closed before reading/writing
subprocess.run(["powershell", "-Command", "Stop-Process -Name POWERPNT -Force -ErrorAction SilentlyContinue"], capture_output=True)

prs = Presentation(pptx_path)
slide = prs.slides[2] # Slide 3 (0-indexed: slide 2)

def update_run_text(shape, paragraph_idx, text):
    p = shape.text_frame.paragraphs[paragraph_idx]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.text = text

# 1. Title
# Shape 7 has 2 paragraphs: [0] "Operational Shift", [1] "Blind Bench vs Cryptographic Proof"
update_run_text(slide.shapes[7], 0, "Operational Shift")
update_run_text(slide.shapes[7], 1, "Traditional CCTV vs Provenance Packing Proof")

# 2. Subtitle
update_run_text(slide.shapes[8], 0, "Replacing hours of manual footage hunting with instant, indisputable packaging video for every parcel.")

# 3. Card Headers
update_run_text(slide.shapes[11], 0, "BEFORE: TRADITIONAL CCTV")
update_run_text(slide.shapes[12], 0, "AFTER: PROVENANCE STATION PROOF")

# 4. Row 1: Search Time
update_run_text(slide.shapes[15], 0, "SEARCH TIME")
update_run_text(slide.shapes[16], 0, "28 min manual scrub across multiple cameras")
update_run_text(slide.shapes[17], 0, "SEARCH TIME")
update_run_text(slide.shapes[18], 0, "Sub-2s instant query by barcode number")

# 5. Row 2: Search Accuracy (was CLOCK DRIFT)
update_run_text(slide.shapes[21], 0, "SEARCH ACCURACY")
update_run_text(slide.shapes[22], 0, "Guesswork based on estimated dispatch time")
update_run_text(slide.shapes[23], 0, "SEARCH ACCURACY")
update_run_text(slide.shapes[24], 0, "Exact clip automatically tagged to tracking ID")

# 6. Row 3: Packing Coverage (was PRE-SCAN FOOTAGE)
update_run_text(slide.shapes[27], 0, "PACKING COVERAGE")
update_run_text(slide.shapes[28], 0, "Misses pre-scan packing steps before label scan")
update_run_text(slide.shapes[29], 0, "PACKING COVERAGE")
update_run_text(slide.shapes[30], 0, "5-second automatic buffer captures item placement")

# 7. Row 4: Video Clarity (was EVIDENCE INTEGRITY)
update_run_text(slide.shapes[33], 0, "VIDEO CLARITY")
update_run_text(slide.shapes[34], 0, "Blurry wide-angle ceiling view with tape glare")
update_run_text(slide.shapes[35], 0, "VIDEO CLARITY")
update_run_text(slide.shapes[36], 0, "1080p Full HD bench view with anti-glare tuning")

# 8. Row 5: Dispute Win Rate
update_run_text(slide.shapes[39], 0, "DISPUTE WIN RATE")
update_run_text(slide.shapes[40], 0, "DISPUTE WIN RATE")
update_run_text(slide.shapes[41], 0, "18%")
update_run_text(slide.shapes[42], 0, "average claim recovery rate")
update_run_text(slide.shapes[43], 0, "82%")
update_run_text(slide.shapes[44], 0, "verified dispute win rate")

# 9. Bottom Banner
update_run_text(slide.shapes[45], 0, "Verified packaging video retrieved in seconds reverses false shortage claims with immediate courier sign-off.")

prs.save(pptx_path)
print("Slide 3 updated successfully in Provenance_Client_Presentation.pptx")
